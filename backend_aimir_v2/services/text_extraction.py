# services/text_extraction.py
from __future__ import annotations

import os
import shutil
import subprocess
from pathlib import Path
from typing import Optional

from pdfminer.high_level import extract_text as pdf_extract_text
from docx import Document as DocxDocument

# гарантируем PATH (важно для macOS / Homebrew)
os.environ.setdefault("PATH", "")
if "/opt/homebrew/bin" not in os.environ["PATH"]:
    os.environ["PATH"] = "/opt/homebrew/bin:" + os.environ["PATH"]


def extract_text_to_file(bin_path: Path, max_bytes: int = 5 * 1024 * 1024) -> Path | None:
    """
    Пишет текст в <doc_dir>/text.txt. Возвращает путь или None.
    Никогда не бросает наружу исключения — безопасно вызывать из роутера после записи файла.
    """
    try:
        doc_dir = bin_path.parent
        doc_dir.mkdir(parents=True, exist_ok=True)
        text = _extract_text(bin_path) or ""
        txt = text.encode("utf-8", errors="ignore")
        if len(txt) > max_bytes:
            txt = txt[:max_bytes]
        out = doc_dir / "text.txt"
        with out.open("wb") as f:
            f.write(txt)
        return out
    except Exception:
        return None


def _extract_text(path: Path) -> Optional[str]:
    ext = path.suffix.lower()
    try:
        if ext == ".pdf":
            t = _extract_text_from_pdf(path)
            _log_backend(path, "pdf(pymupdf/pdfminer)" if t.strip() else "pdf(empty)")
            return t
        if ext == ".docx":
            t = _extract_text_from_docx(path)
            _log_backend(path, "python-docx" if t.strip() else "python-docx(empty)")
            return t
        if ext == ".doc":
            t = _extract_text_from_doc(path)
            _log_backend(path, "doc-pipeline" if t.strip() else "doc-pipeline(empty)")
            return t
        return ""
    except Exception:
        return ""


def _extract_text_from_pdf(path: Path) -> str:
    # Основной движок — PyMuPDF (fitz): сохраняет корректный порядок чтения.
    # pdfminer на ряде PDF (двухколоночная верстка, выравнивание по ширине)
    # перемешивает текст, из-за чего пункты выглядят «обрезанными».
    try:
        import fitz  # PyMuPDF
        with fitz.open(str(path)) as doc:
            text = "".join(page.get_text("text") for page in doc)
        if text.strip():
            return text.strip()
    except Exception:
        pass
    # Фолбэк — pdfminer (если PyMuPDF недоступен или не смог прочитать файл).
    return (pdf_extract_text(str(path)) or "").strip()


def _extract_text_from_docx(path: Path) -> str:
    doc = DocxDocument(str(path))
    parts: list[str] = []
    for p in doc.paragraphs:
        t = (p.text or "").strip("\x00")
        if t:
            parts.append(t)
    # таблицы
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip("\x00") for cell in row.cells]
            if any(cells):
                parts.append("\t".join(cells))
    return "\n".join(parts).strip()


# ---------- .DOC: несколько стратегий ----------
def _extract_text_from_doc(path: Path) -> str:
    # 1) LibreOffice (soffice) -> .docx
    conv = _convert_with_soffice(path, to="docx")
    if conv and conv.exists():
        try:
            text = _extract_text_from_docx(conv)
            _safe_unlink(conv)
            if text.strip():
                return text
        except Exception:
            _safe_unlink(conv)

    # 2) unoconv -> txt
    txt = _convert_with_unoconv(path, to="txt")
    if txt and txt.exists():
        try:
            data = txt.read_text(encoding="utf-8", errors="ignore")
            _safe_unlink(txt)
            if data.strip():
                return data
        except Exception:
            _safe_unlink(txt)

    # 3) antiword (stdout)
    text = _run_tool_capture(["antiword", str(path)])
    if text.strip():
        return text

    # 4) catdoc (stdout)
    text = _run_tool_capture(["catdoc", "-w", str(path)])
    if text.strip():
        return text

    # 5) textract (если установлен)
    try:
        import textract  # type: ignore
        raw = textract.process(str(path))
        text = (raw or b"").decode("utf-8", errors="ignore")
        if text.strip():
            return text
    except Exception:
        pass

    # 6) без инструментов — возвращаем пусто
    return ""


# ---------- helpers ----------
def _convert_with_soffice(src: Path, to: str) -> Optional[Path]:
    soffice = shutil.which("soffice")
    if not soffice:
        return None
    out_dir = src.parent
    try:
        subprocess.run(
            [soffice, "--headless", "--convert-to", to, "--outdir", str(out_dir), str(src)],
            check=True, timeout=30, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL
        )
        candidate = out_dir / f"{src.stem}.{to}"
        return candidate if candidate.exists() else None
    except Exception:
        return None


def _convert_with_unoconv(src: Path, to: str) -> Optional[Path]:
    unoconv = shutil.which("unoconv")
    if not unoconv:
        return None
    out = src.parent / f"{src.stem}.{to}"
    try:
        subprocess.run(
            [unoconv, "-f", to, "-o", str(out), str(src)],
            check=True, timeout=30, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL
        )
        return out if out.exists() else None
    except Exception:
        return None


def _run_tool_capture(cmd: list[str]) -> str:
    if not shutil.which(cmd[0]):
        return ""
    try:
        res = subprocess.run(cmd, timeout=20, check=False, stdout=subprocess.PIPE, stderr=subprocess.DEVNULL)
        return (res.stdout or b"").decode("utf-8", errors="ignore")
    except Exception:
        return ""


def _safe_unlink(p: Path) -> None:
    try:
        p.unlink()
    except Exception:
        pass


def _log_backend(path: Path, backend: str) -> None:
    # мягкий лог в stdout — видно в консоли uvicorn/pycharm
    try:
        print(f"[text-extract] {path.suffix.lower()} via {backend}")
    except Exception:
        pass
